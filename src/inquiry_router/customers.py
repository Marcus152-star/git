from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

from openpyxl import load_workbook

from .models import CustomerMatch, CustomerRecord


EMAIL_RE = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.IGNORECASE)

PUBLIC_EMAIL_DOMAINS = {
    "126.com",
    "163.com",
    "aol.com",
    "gmail.com",
    "googlemail.com",
    "hotmail.com",
    "icloud.com",
    "live.com",
    "mail.ru",
    "outlook.com",
    "proton.me",
    "protonmail.com",
    "qq.com",
    "yahoo.com",
    "yahoo.co.jp",
    "yandex.com",
    "yandex.ru",
}

REQUIRED_COLUMNS = {
    "客户编号",
    "公司名称",
    "国家",
    "公司网址",
    "公司邮箱",
    "联系人工作邮箱",
    "联系人个人邮箱",
    "当前跟进人",
}

EMAIL_COLUMNS = ("公司邮箱", "联系人工作邮箱", "联系人个人邮箱")


def normalize_email(value: str) -> str:
    return value.strip().strip("<>[](){}.,;:\"'").lower()


def extract_emails(value: object) -> list[str]:
    return [normalize_email(match) for match in EMAIL_RE.findall(str(value or ""))]


def email_domain(email: str) -> str:
    return email.rsplit("@", 1)[1].lower() if "@" in email else ""


def normalize_company(value: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", value.casefold())


def website_domain(value: object) -> str:
    raw = str(value or "").strip().lower()
    if not raw:
        return ""
    parsed = urlparse(raw if "://" in raw else f"https://{raw}")
    host = (parsed.hostname or "").removeprefix("www.")
    return host


@dataclass
class CustomerIndex:
    records: list[CustomerRecord]
    by_email: dict[str, list[CustomerRecord]]
    by_domain: dict[str, list[CustomerRecord]]
    by_company: dict[str, list[CustomerRecord]]

    @classmethod
    def from_xlsx(cls, path: str | Path) -> "CustomerIndex":
        return cls.from_xlsx_files([path])

    @classmethod
    def from_xlsx_files(cls, paths: Iterable[str | Path]) -> "CustomerIndex":
        resolved_paths = [Path(path) for path in paths]
        if not resolved_paths:
            raise ValueError("At least one customer workbook is required")

        records: list[CustomerRecord] = []
        for path in resolved_paths:
            records.extend(cls._read_records(path))

        by_email: dict[str, list[CustomerRecord]] = defaultdict(list)
        by_domain: dict[str, list[CustomerRecord]] = defaultdict(list)
        by_company: dict[str, list[CustomerRecord]] = defaultdict(list)
        for record in records:
            for email in record.emails:
                by_email[email].append(record)
            for domain in record.domains:
                by_domain[domain].append(record)
            company_key = normalize_company(record.company_name)
            if company_key:
                by_company[company_key].append(record)

        return cls(records, dict(by_email), dict(by_domain), dict(by_company))

    @staticmethod
    def _read_records(path: Path) -> list[CustomerRecord]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [str(value or "").strip() for value in next(rows)]
            except StopIteration as exc:
                raise ValueError(f"Customer workbook is empty: {path.name}") from exc

            missing = sorted(REQUIRED_COLUMNS.difference(headers))
            if missing:
                raise ValueError(
                    f"Customer workbook {path.name} is missing columns: {', '.join(missing)}"
                )

            records: list[CustomerRecord] = []
            for row_number, values in enumerate(rows, start=2):
                data = dict(zip(headers, values))
                if not any(value not in (None, "") for value in values):
                    continue

                emails = sorted(
                    {
                        email
                        for column in EMAIL_COLUMNS
                        for email in extract_emails(data.get(column))
                    }
                )
                domains = {
                    domain
                    for domain in (email_domain(email) for email in emails)
                    if domain and domain not in PUBLIC_EMAIL_DOMAINS
                }
                site_domain = website_domain(data.get("公司网址"))
                if site_domain and site_domain not in PUBLIC_EMAIL_DOMAINS:
                    domains.add(site_domain)

                records.append(
                    CustomerRecord(
                        row_number=row_number,
                        customer_id=str(data.get("客户编号") or "").strip(),
                        company_name=str(data.get("公司名称") or "").strip(),
                        country=str(data.get("国家") or "").strip(),
                        owner=str(data.get("当前跟进人") or "").strip(),
                        emails=tuple(emails),
                        domains=tuple(sorted(domains)),
                        source_path=str(path.resolve()),
                    )
                )
            return records
        finally:
            workbook.close()

    @staticmethod
    def _result(records: list[CustomerRecord], reason: str, matched_email: str = "") -> CustomerMatch:
        unique = {
            (record.customer_id, record.source_path, record.row_number): record
            for record in records
        }
        values = list(unique.values())
        owners = {record.owner for record in values if record.owner}
        countries = tuple(sorted({record.country for record in values if record.country}))
        companies = tuple(sorted({record.company_name for record in values if record.company_name}))
        customer_ids = tuple(sorted({record.customer_id for record in values if record.customer_id}))
        if len(owners) == 1:
            return CustomerMatch(
                status="matched",
                reason=reason,
                matched_email=matched_email,
                owner=next(iter(owners)),
                customer_ids=customer_ids,
                company_names=companies,
                countries=countries,
            )
        if not owners:
            return CustomerMatch(
                status="ambiguous",
                reason=f"{reason}; matched record has no current owner",
                matched_email=matched_email,
                customer_ids=customer_ids,
                company_names=companies,
                countries=countries,
            )
        return CustomerMatch(
            status="ambiguous",
            reason=f"{reason}; records have multiple current owners",
            matched_email=matched_email,
            customer_ids=customer_ids,
            company_names=companies,
            countries=countries,
        )

    def match(self, emails: list[str], company_name: str = "") -> CustomerMatch:
        normalized_emails = [normalize_email(email) for email in emails]
        for email in normalized_emails:
            records = self.by_email.get(email, [])
            if records:
                return self._result(records, "exact email match", email)

        for email in normalized_emails:
            domain = email_domain(email)
            if not domain or domain in PUBLIC_EMAIL_DOMAINS:
                continue
            records = self.by_domain.get(domain, [])
            if records:
                return self._result(records, "unique business domain match", email)

        company_key = normalize_company(company_name)
        if company_key and company_key in self.by_company:
            return self._result(self.by_company[company_key], "exact normalized company-name match")

        return CustomerMatch(status="not_found", reason="no exact email, business domain, or company match")

    def summary(self) -> dict:
        owner_counts = Counter(record.owner or "<空>" for record in self.records)
        country_counts = Counter(record.country or "<空>" for record in self.records)
        all_emails = [email for record in self.records for email in record.emails]
        duplicate_values = sum(count - 1 for count in Counter(all_emails).values() if count > 1)
        source_files = sorted({record.source_path for record in self.records if record.source_path})
        return {
            "source_file_count": len(source_files),
            "source_files": source_files,
            "records": len(self.records),
            "records_with_current_owner": sum(bool(record.owner) for record in self.records),
            "records_with_country": sum(bool(record.country) for record in self.records),
            "records_with_email": sum(bool(record.emails) for record in self.records),
            "records_without_email": sum(not record.emails for record in self.records),
            "unique_emails": len(set(all_emails)),
            "unique_business_domains": len(self.by_domain),
            "duplicate_email_values": duplicate_values,
            "owner_counts": dict(owner_counts.most_common()),
            "top_countries": dict(country_counts.most_common(15)),
        }

